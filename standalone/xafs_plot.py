#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xafs_plot.py —— Athena 工程文件(.prj) XAFS 一键分析与作图（不需要 Claude）

用法：
    python xafs_plot.py athena.prj
    python xafs_plot.py athena.prj --sample-name "FeCoNiRu-Cu"
    python xafs_plot.py athena.prj --config athena_config.json

第一次运行会在 .prj 旁边生成 <文件名>_config.json（自动识别的元素、样品/标样、
显示名、颜色、k 范围、阴影带等）。修改它之后再运行一次即可按你的设置重画。

输出（在 <文件名>_xafs_output 文件夹）：
    <元素>_XAFS.png / .pdf     论文风格图：左 FT-EXAFS（k³，堆叠），右 XANES
    XAFS_origin_plot.py       Origin 2021+ 内置 Python 作图脚本（数据已嵌入）
    XAFS_data.xlsx            作图数据（可直接导入 Origin）
    report.txt                数据检查与分析报告（标定、价态、LCF、FT 峰位等）
"""
import argparse
import json
import math
import os
import re
import sys
import warnings

import numpy as np

warnings.filterwarnings('ignore')

try:
    from larch.io import read_athena
    from larch.xafs import pre_edge, autobk, xftf
except Exception:
    sys.exit('需要先安装 xraylarch：pip install xraylarch\n'
             '（若安装失败，见 README 中的“安装问题”一节）')

from scipy.optimize import nnls
from scipy.signal import find_peaks

HERE = os.path.dirname(os.path.abspath(__file__))

# ------------------------------------------------------------------
# 吸收边表（eV）：K 边，以及常见 L3 边（键名带 _L3）
# ------------------------------------------------------------------
EDGES = {
    'Ti': 4966, 'V': 5465, 'Cr': 5989, 'Mn': 6539, 'Fe': 7112, 'Co': 7709,
    'Ni': 8333, 'Cu': 8979, 'Zn': 9659, 'Ga': 10367, 'Ge': 11103, 'As': 11867,
    'Se': 12658, 'Br': 13474, 'Sr': 16105, 'Y': 17038, 'Zr': 17998, 'Nb': 18986,
    'Mo': 20000, 'Ru': 22117, 'Rh': 23220, 'Pd': 24350, 'Ag': 25514, 'Cd': 26711,
    'In': 27940, 'Sn': 29200, 'Sb': 30491, 'Ce': 40443,
    'W_L3': 10207, 'Re_L3': 10535, 'Os_L3': 10871, 'Ir_L3': 11215,
    'Pt_L3': 11564, 'Au_L3': 11919, 'Ce_L3': 5723, 'La_L3': 5483,
}

SAMPLE_COLORS = ['#D73027', '#A50026', '#F46D43', '#7F0000']
REF_COLORS = ['#FC8D59', '#91BFDB', '#FDAE61', '#74ADD1', '#878787']
FOIL_COLOR = '#4575B4'
BAND_MO = '#FEE090'      # M–O 壳层阴影（浅黄）
BAND_MM = '#C8DFED'      # M–M 壳层阴影（浅蓝）

DEFAULT_STYLE = {
    'font': 'Arial',
    'origin_font_size': 36,     # Origin 脚本里所有文字字号
    'origin_frame_width': 3,    # Origin 边框线宽
    'origin_line_width': 3,     # Origin 曲线线宽
    'figure_font_size': 22,     # matplotlib 图的字号
}


# ------------------------------------------------------------------
# 工具函数
# ------------------------------------------------------------------
def edge_symbol(key):
    return key.split('_')[0]


def edge_name(key):
    return f'{edge_symbol(key)} L$_3$-edge' if key.endswith('_L3') else f'{key} K-edge'


def edge_name_plain(key):
    return f'{edge_symbol(key)} L3-edge' if key.endswith('_L3') else f'{key} K-edge'


def detect_edge(e0, emin, emax):
    best, bd = None, 1e9
    for k, v in EDGES.items():
        d = abs(e0 - v)
        if d < bd:
            best, bd = k, d
    if bd <= 80:
        return best
    # e0 不可靠时，用数据范围中间段再找一次
    for k, v in EDGES.items():
        if emin + 30 < v < emax - 50:
            return k
    return None


def fix_case(name, sym):
    if name.upper().startswith(sym.upper()) and not name.startswith(sym):
        name = sym + name[len(sym):]
    return name


def clean_name(label, sym):
    n = label
    n = re.sub(r'[-_ ]?(default|stand|std|standard|ref)[-_ ]?\d*', '', n, flags=re.I)
    n = re.sub(r'[-_ ]\d+$', '', n)
    n = re.sub(r'[-_\s]*foil', ' foil', n, flags=re.I).strip()
    n = re.sub(r'^post-', '', n, flags=re.I)
    n = fix_case(n, sym)
    return n.strip(' -_') or label


FORMULA_TOKEN = re.compile(r'([A-Z][a-z]?)(\d*)|\((OH)\)(\d*)')


def parse_formula(name):
    """返回 {元素: 数量}；不是化学式则返回 None"""
    if not re.fullmatch(r'(?:[A-Z][a-z]?\d*|\(OH\)\d*)+', name):
        return None
    comp = {}
    for m in FORMULA_TOKEN.finditer(name):
        if m.group(1):
            comp[m.group(1)] = comp.get(m.group(1), 0) + int(m.group(2) or 1)
        else:
            k = int(m.group(4) or 1)
            comp['O'] = comp.get('O', 0) + k
            comp['H'] = comp.get('H', 0) + k
    return comp


def oxidation_state(name, sym, role):
    if role == 'foil':
        return 0.0
    comp = parse_formula(name)
    if not comp or sym not in comp or 'O' not in comp:
        return None
    others = [e for e in comp if e not in (sym, 'O', 'H')]
    if others:
        return None
    return (2 * comp['O'] - comp.get('H', 0)) / comp[sym]


def guess_role(label, name, sym):
    if 'foil' in label.lower():
        return 'foil'
    if re.search(r'default|stand|std|ref|标样', label, re.I):
        return 'reference'
    comp = parse_formula(name)
    if comp and 'O' in comp and sym in comp:
        return 'reference'
    return 'sample'


def origin_label(name):
    """化学式数字转 Origin 下标：Fe2O3 -> Fe\\-(2)O\\-(3)"""
    return re.sub(r'(?<=[A-Za-z)])(\d+)', r'\\-(\1)', name)


def mpl_label(name):
    return re.sub(r'(?<=[A-Za-z)])(\d+)', r'$_{\1}$', name)


# ------------------------------------------------------------------
# 读取 & 处理
# ------------------------------------------------------------------
def load_groups(prj):
    p = read_athena(prj, do_preedge=False, do_bkg=False, do_fft=False)
    groups = []
    for k in dir(p):
        g = getattr(p, k)
        if not hasattr(g, 'athena_params'):
            continue
        e, mu = np.asarray(g.energy, float), np.asarray(g.mu, float)
        idx = np.argsort(e)
        e, mu = e[idx], mu[idx]
        keep = np.r_[True, np.diff(e) > 0.05]
        g.n_dup = int((~keep).sum())
        shift = float(getattr(g.athena_params.bkg, 'eshift', 0) or 0)
        g.energy, g.mu = e[keep] + shift, mu[keep]
        g.eshift = shift
        g.athena_e0 = float(g.athena_params.bkg.e0)
        groups.append(g)
    return groups


def process(g, kmin, kmax, kweight):
    e0 = g.athena_e0
    emin, emax = g.energy.min(), g.energy.max()
    pre1 = max(-150, emin - e0 + 5)
    pre2 = min(-30, pre1 + 40)
    span = emax - e0
    norm2 = min(400, span - 10)
    norm1 = min(150, norm2 * 0.4)
    pre_edge(g, e0=e0, pre1=pre1, pre2=pre2, norm1=norm1, norm2=norm2, nnorm=2)
    autobk(g, e0=e0, rbkg=1.0, kweight=kweight)
    xftf(g, kmin=kmin, kmax=kmax, dk=1, kweight=kweight, window='hanning')


def edge_metrics(g, edge_e):
    e, n = g.energy, g.norm
    m = (e > edge_e - 30) & (e < edge_e + 40)
    d = np.gradient(n, e)
    emaxd = float(e[m][np.argmax(d[m])]) if m.any() else float('nan')
    mm = (e > edge_e - 30) & (e < edge_e + 60)
    ehalf = float('nan')
    ee, nn = e[mm], n[mm]
    idx = np.where(nn >= 0.5)[0]
    if len(idx) and idx[0] > 0:
        i = idx[0]
        ehalf = float(np.interp(0.5, [nn[i - 1], nn[i]], [ee[i - 1], ee[i]]))
    w = (e > edge_e) & (e < edge_e + 40)
    wl = float(g.flat[w].max()) if w.any() else float('nan')
    return emaxd, ehalf, wl


def ft_peaks(g, rmin=0.8, rmax=4.0):
    m = (g.r > rmin) & (g.r < rmax)
    y = g.chir_mag[m]
    if not len(y):
        return []
    pk, _ = find_peaks(y, prominence=0.08 * y.max())
    return [(float(g.r[m][i]), float(y[i] / y.max())) for i in pk]


# ------------------------------------------------------------------
# 配置
# ------------------------------------------------------------------
def build_default_config(groups, sample_name=None):
    spectra, elements = [], {}
    for g in groups:
        key = detect_edge(g.athena_e0, g.energy.min(), g.energy.max())
        if key is None:
            spectra.append(dict(label=g.label, element=None, role='skip', name=g.label,
                                color='#000000', include=False,
                                note='无法识别吸收边，已跳过'))
            continue
        sym = edge_symbol(key)
        nm = clean_name(g.label, sym)
        role = guess_role(g.label, nm, sym)
        if role == 'sample':
            nm = sample_name or g.label
        include = abs(g.eshift) <= 20
        note = '' if include else f'Athena 中能量平移 {g.eshift:.1f} eV 过大，默认不画，请核实'
        spectra.append(dict(label=g.label, element=key, role=role, name=nm,
                            color='', include=include, note=note))
        elements.setdefault(key, {})
    # 排序与配色：每个元素 —— 氧化物（价态高→低）、箔、样品
    out = []
    for key in elements:
        sym = edge_symbol(key)
        sp = [s for s in spectra if s['element'] == key]
        refs = [s for s in sp if s['role'] == 'reference']
        refs.sort(key=lambda s: -(oxidation_state(s['name'], sym, 'reference') or -1))
        foils = [s for s in sp if s['role'] == 'foil']
        samples = [s for s in sp if s['role'] == 'sample']
        inc = [s for s in refs if s['include']] + [s for s in refs if not s['include']]
        for i, s in enumerate(sorted(inc, key=lambda s: (not s['include'], oxidation_state(s['name'], sym, 'reference') or 99))):
            s['color'] = REF_COLORS[i % len(REF_COLORS)]
        for s in foils:
            s['color'] = FOIL_COLOR
        for i, s in enumerate(samples):
            s['color'] = SAMPLE_COLORS[i % len(SAMPLE_COLORS)]
        out += refs + foils + samples
    out += [s for s in spectra if s['element'] is None]
    for key in elements:
        e = EDGES[key]
        x0 = int(math.floor((e - 5) / 10.0) * 10)
        elements[key] = dict(kmin=3, kmax='auto', kweight=3,
                             xanes_range=[x0, x0 + 50], bands='auto')
    return dict(
        说明='role: sample/reference/foil；include=false 不画；列表顺序 = FT 图从上到下的顺序；'
           'kmax 可填数字；bands 可填 [[1.5,"#FEE090"],[2.7,"#C8DFED"]] 或 []；改完重新运行即可',
        style=dict(DEFAULT_STYLE), elements=elements, spectra=out)


# ------------------------------------------------------------------
# FT 堆叠排版（与 Origin 脚本中算法相同）：保证标注不压曲线
# ------------------------------------------------------------------
def ft_layout(R, ys, labels, layer_w_pt, layer_h_pt, font_pt, char_w=0.62, xr=6.0):
    n = len(ys)
    txt_h = font_pt * 1.25
    pt_per_A = layer_w_pt / xr
    order = list(range(n - 1, -1, -1))
    yrange = 1.6 * n
    top = yrange
    for _ in range(30):
        u = yrange / layer_h_pt
        h, gap = txt_h * u, 8 * u
        offs, labs = {}, {}
        off = 0.0
        for j, i in enumerate(order):
            wA = (len(re.sub(r'\$_\{(\d+)\}\$', r'\1', labels[i])) * char_w * font_pt + 12) / pt_per_A
            x0 = xr - 0.12 - wA
            span = [y for x, y in zip(R, ys[i]) if x >= x0]
            m = max(span) if span else 0.3
            if j > 0:
                prev = order[j - 1]
                mn = min(span) if span else 0.0
                off = max(offs[prev] + 1.08, labs[prev][1] + h + gap - mn)
            offs[i] = off
            labs[i] = (xr - 0.12, off + m + gap)
        top = max(offs[order[-1]] + 1.05, labs[order[-1]][1] + h + gap)
        if abs(top + 0.05 - yrange) < 1e-3:
            break
        yrange = top + 0.05
    return offs, labs, -0.05, top


# ------------------------------------------------------------------
# 主流程
# ------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description='Athena .prj XAFS 一键分析与作图')
    ap.add_argument('prj', help='Athena 工程文件 (.prj)')
    ap.add_argument('--config', help='配置文件（默认 <文件名>_config.json，不存在则自动生成）')
    ap.add_argument('--sample-name', help='把所有样品统一显示为这个名字（如 FeCoNiRu-Cu）')
    ap.add_argument('--out', help='输出文件夹（默认 <文件名>_xafs_output）')
    args = ap.parse_args()

    prj = os.path.abspath(args.prj)
    stem = os.path.splitext(os.path.basename(prj))[0]
    base = os.path.dirname(prj)
    cfg_path = os.path.abspath(args.config) if args.config else os.path.join(base, f'{stem}_config.json')
    out_dir = os.path.abspath(args.out) if args.out else os.path.join(base, f'{stem}_xafs_output')
    os.makedirs(out_dir, exist_ok=True)

    print(f'读取 {prj} ...')
    groups = load_groups(prj)
    G = {g.label: g for g in groups}
    print(f'  共 {len(groups)} 条谱')

    if os.path.exists(cfg_path):
        cfg = json.load(open(cfg_path, encoding='utf-8'))
        print(f'使用配置文件：{cfg_path}')
        if args.sample_name:
            for s in cfg['spectra']:
                if s['role'] == 'sample':
                    s['name'] = args.sample_name
    else:
        cfg = build_default_config(groups, args.sample_name)
        json.dump(cfg, open(cfg_path, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
        print(f'已生成配置文件：{cfg_path}\n  （检查样品/标样识别和显示名，修改后重新运行即可）')
    style = dict(DEFAULT_STYLE)
    style.update(cfg.get('style', {}))

    report = [f'XAFS 分析报告  —  {os.path.basename(prj)}', '=' * 60, '']
    skipped = [s for s in cfg['spectra'] if not s.get('include', True)]
    if skipped:
        report.append('未参与作图的谱（可在配置文件中把 include 改为 true）：')
        for s in skipped:
            report.append(f'   ⚠ {s["label"]}：{s.get("note") or "include=false"}')
        report.append('')
    origin_data, excel_sheets = {}, []

    for key, ecfg in cfg['elements'].items():
        sym = edge_symbol(key)
        edge_e = EDGES[key]
        sp = [s for s in cfg['spectra'] if s['element'] == key and s.get('include', True)
              and s['label'] in G]
        if not sp:
            continue
        samples = [s for s in sp if s['role'] == 'sample']
        refs = [s for s in sp if s['role'] != 'sample']
        report += [f'【{edge_name_plain(key)}】（标准边位置 {edge_e} eV）', '']

        # k 范围
        kweight = ecfg.get('kweight', 3)
        kmin = ecfg.get('kmin', 3)
        kmax = ecfg.get('kmax', 'auto')
        if kmax == 'auto':
            kmax = 10.0
            for s in samples or sp:
                g = G[s['label']]
                kav = math.sqrt(max(g.energy.max() - g.athena_e0, 1) / 3.81)
                kmax = min(kmax, kav - 0.5)
            # 先粗处理一次得到边跳，弱信号缩短 k 范围
            for s in samples:
                g = G[s['label']]
                pre_edge(g, e0=g.athena_e0)
                if g.edge_step < 0.02:
                    kmax = min(kmax, 8.0)
            kmax = round(kmax * 2) / 2
        for s in sp:
            process(G[s['label']], kmin, kmax, kweight)

        # ---- 数据检查 ----
        report.append('  数据检查：')
        issues = 0
        for s in sp:
            g = G[s['label']]
            emaxd, ehalf, wl = edge_metrics(g, edge_e)
            s['_m'] = (emaxd, ehalf, wl)
            if s['role'] == 'foil':
                # 金属箔：取吸收边起点附近的第一个导数峰
                e, d = g.energy, np.gradient(g.norm, g.energy)
                m = (e > edge_e - 25) & (e < edge_e + 25)
                pk, _ = find_peaks(d[m], prominence=0.1 * d[m].max())
                e_first = float(e[m][pk[0]]) if len(pk) else emaxd
                dev = e_first - edge_e
                flag = '⚠ ' if abs(dev) > 1.5 else '✓ '
                if abs(dev) > 1.5:
                    issues += 1
                report.append(f'   {flag}{s["name"]}：一阶导第一个峰 {e_first:.1f} eV，'
                              f'与标准值相差 {dev:+.1f} eV'
                              + ('（同批测试的样品标定可能需要校正）' if abs(dev) > 1.5 else ''))
            if abs(g.eshift) > 5:
                issues += 1
                report.append(f'   ⚠ {s["name"]}：Athena 中设置了 {g.eshift:+.1f} eV 能量平移，请核实是否合理')
            if g.n_dup:
                report.append(f'   · {s["name"]}：删除了 {g.n_dup} 个重复能量点')
            if s['role'] == 'sample' and g.edge_step < 0.02:
                issues += 1
                report.append(f'   ⚠ {s["name"]}：边跳很小（{g.edge_step:.4f}），信号弱，'
                              f'EXAFS 只用到 k = {kmax} Å⁻¹')
        foils = [s for s in refs if s['role'] == 'foil']
        for s in samples:
            for f in foils:
                gs, gf = G[s['label']], G[f['label']]
                x = np.arange(edge_e - 20, edge_e + 80, 0.5)
                c = np.corrcoef(np.interp(x, gs.energy, gs.norm), np.interp(x, gf.energy, gf.norm))[0, 1]
                if c > 0.998:
                    issues += 1
                    report.append(f'   ⚠ {s["name"]} 与 {f["name"]} 的 XANES 几乎重合（相关系数 {c:.4f}）。'
                                  f'若样品长在含 {sym} 的基底上并连基底一起测，信号可能主要来自基底。')
        if not issues:
            report.append('   ✓ 未发现明显问题')
        report.append('')

        # ---- 边位置 / 白线 ----
        report.append('  吸收边与白线（E_max-deriv：导数最大处；E_0.5：归一化 0.5 处；WL：白线强度）：')
        for s in refs + samples:
            emaxd, ehalf, wl = s['_m']
            report.append(f'   {s["name"]:<22s} E_max-deriv = {emaxd:8.1f}   E_0.5 = {ehalf:8.1f}   WL = {wl:.2f}')
        # 价态插值
        pts = [(oxidation_state(s['name'], sym, s['role']), s['_m'][1]) for s in refs]
        pts = [(o, e) for o, e in pts if o is not None and not math.isnan(e) and o > 0]
        if len({o for o, _ in pts}) >= 2:
            o_, e_ = np.array(pts).T
            a, b = np.polyfit(e_, o_, 1)
            for s in samples:
                est = a * s['_m'][1] + b
                if not (e_.min() - 0.3 <= s['_m'][1] <= e_.max() + 0.3) or (e_.max() - e_.min()) < 1.5:
                    report.append(f'   → {s["name"]}：E_0.5 不在氧化物标样之间（或标样间差距太小），'
                                  f'不做价态插值；请结合白线形状和 LCF 判断')
                    continue
                report.append(f'   → {s["name"]}：按氧化物标样 E_0.5 线性插值，平均价态 ≈ {est:+.2f}'
                              f'（标样与样品若非同批测试，能量误差约 ±1 eV，仅供参考）')
        report.append('')

        # ---- LCF ----
        if len(refs) >= 1:
            report.append(f'  线性组合拟合 LCF（{edge_e - 20} ~ {edge_e + 40} eV，非负最小二乘）：')
            x = np.arange(edge_e - 20, edge_e + 40, 0.25)
            A = np.array([np.interp(x, G[r['label']].energy, G[r['label']].norm) for r in refs]).T
            for s in samples:
                y = np.interp(x, G[s['label']].energy, G[s['label']].norm)
                w, _ = nnls(A, y)
                R = np.sum((A @ w - y) ** 2) / np.sum(y ** 2)
                parts = '，'.join(f'{r["name"]} {wi * 100:.0f}%' for r, wi in zip(refs, w) if wi > 0.005)
                report.append(f'   {s["name"]}：{parts}（权重和 {w.sum():.2f}，R = {R:.4f}'
                              + ('，拟合较差，标样可能不能代表样品' if R > 0.01 else '') + '）')
            report.append('')

        # ---- FT 峰 ----
        report.append(f'  FT-EXAFS 峰位（k{kweight} 加权，k = {kmin}–{kmax} Å⁻¹，未相位校正；括号内为相对高度）：')
        for s in refs + samples:
            pk = ft_peaks(G[s['label']])
            report.append(f'   {s["name"]:<22s} ' + '  '.join(f'{r:.2f} Å({h:.2f})' for r, h in pk))
        for s in samples:
            for f in foils:
                ratio = G[s['label']].chir_mag.max() / G[f['label']].chir_mag.max()
                report.append(f'   {s["name"]} / {f["name"]} 最高峰强度比 = {ratio:.2f}')
        report.append('')

        # ---- 阴影带 ----
        bands = ecfg.get('bands', 'auto')
        if bands == 'auto':
            bands = []
            if samples:
                pk = sorted(ft_peaks(G[samples[0]['label']], 0.9, 3.4), key=lambda t: -t[1])[:2]
                metallic = False
                if foils and pk:
                    fpk = sorted(ft_peaks(G[foils[0]['label']], 1.8, 2.8), key=lambda t: -t[1])
                    metallic = bool(fpk) and abs(pk[0][0] - fpk[0][0]) < 0.15
                for r, h in sorted(pk):
                    if h < 0.25:
                        continue
                    col = BAND_MO if (r < 2.0 and not metallic) else BAND_MM
                    bands.append([round(r, 2), col])
        ecfg['_bands'] = bands

        # ---- 收集作图数据 ----
        order = [s for s in sp]                     # 配置顺序 = FT 从上到下
        R = G[order[0]['label']].r
        mR = R <= 6
        curves = []
        for s in order:
            g = G[s['label']]
            y = g.chir_mag[mR] / g.chir_mag[mR & (g.r > 0.5)].max()
            x0, x1 = ecfg['xanes_range']
            mm = (g.energy >= x0 - 30) & (g.energy <= x1 + 100)
            curves.append(dict(name=s['name'], label=origin_label(s['name']), mpl=mpl_label(s['name']),
                               color=s['color'], offset=0.0, role=s['role'],
                               ft=[round(float(v), 5) for v in y],
                               ft_raw=[round(float(v), 5) for v in g.chir_mag[mR]],
                               E=[round(float(v), 3) for v in g.energy[mm]],
                               mu=[round(float(v), 5) for v in g.flat[mm]]))
        xan_order = [c['name'] for c in curves if c['role'] == 'sample'] + \
                    [c['name'] for c in curves if c['role'] != 'sample']
        x0, x1 = ecfg['xanes_range']
        xstep = 20 if (x1 - x0) >= 40 else 10
        origin_data[key] = dict(R=[round(float(v), 4) for v in R[mR]], curves=curves, km=kmax,
                                xr=[x0, x1], xstep=xstep,
                                bands=[b[0] for b in bands], band_colors=[b[1] for b in bands],
                                xan_order=xan_order)

        # ---- matplotlib 图 ----
        draw_figure(key, origin_data[key], style, kweight, kmin, kmax, out_dir)
        excel_sheets.append((key, origin_data[key]))
        print(f'  {edge_name_plain(key)}：完成')

    # ---- 输出 ----
    open(os.path.join(out_dir, 'report.txt'), 'w', encoding='utf-8').write('\n'.join(report))
    write_origin(origin_data, style, out_dir)
    write_excel(excel_sheets, out_dir)
    json.dump(cfg_strip(cfg), open(cfg_path, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    print(f'\n全部完成，结果在：{out_dir}')
    print('  report.txt 里有数据检查结果（⚠ 开头的条目请重点核实）')


def cfg_strip(cfg):
    c = json.loads(json.dumps(cfg, default=str))
    for s in c['spectra']:
        s.pop('_m', None)
    for e in c['elements'].values():
        e.pop('_bands', None)
    return c


def draw_figure(key, d, style, kweight, kmin, kmax, out_dir):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    fams = {f.name for f in font_manager.fontManager.ttflist}
    font = next((f for f in [style['font'], 'Arial', 'Liberation Sans', 'Helvetica'] if f in fams), 'DejaVu Sans')
    fs = style['figure_font_size']
    plt.rcParams.update({'font.family': font, 'font.size': fs, 'axes.linewidth': 1.8,
                         'xtick.direction': 'out', 'xtick.major.width': 1.8, 'xtick.major.size': 7,
                         'xtick.minor.visible': True, 'xtick.minor.size': 4, 'mathtext.default': 'regular',
                         'savefig.bbox': 'tight'})
    W, H = 15, 6.4
    f, (a, b) = plt.subplots(1, 2, figsize=(W, H), gridspec_kw={'wspace': 0.22})
    lw_pt = W * 72 * 0.36
    lh_pt = H * 72 * 0.77
    R = d['R']
    ys = [c['ft'] for c in d['curves']]
    offs, labs, ylo, yhi = ft_layout(R, ys, [c['mpl'] for c in d['curves']], lw_pt, lh_pt, fs * 0.95)
    for bc, col in zip(d['bands'], d['band_colors']):
        a.axvspan(bc - 0.2, bc + 0.2, color=col, lw=0, zorder=0)
    for i, c in enumerate(d['curves']):
        a.plot(R, [v + offs[i] for v in c['ft']], color=c['color'], lw=2.6)
        x, y = labs[i]
        a.text(x, y, c['mpl'], ha='right', va='bottom', fontsize=fs * 0.95)
    a.set_xlim(0, 6)
    a.set_xticks([0, 2, 4, 6])
    a.set_ylim(ylo, yhi)
    a.set_yticks([])
    a.tick_params(top=False, right=False)
    a.set_xlabel('R (Å)')
    a.set_ylabel(f'FT$\\it{{k}}^{{{kweight}}}\\chi(\\it{{k}})$ (a.u.)')
    a.set_title(f'{edge_name(key)}, $\\it{{k}}$ = {kmin}–{kmax} Å$^{{-1}}$', fontsize=fs * 0.7, loc='left')
    cv = {c['name']: c for c in d['curves']}
    x0, x1 = d['xr']
    for nm in d['xan_order']:
        c = cv[nm]
        b.plot(c['E'], c['mu'], color=c['color'], lw=2.6, label=c['mpl'])
    vals = [v for nm in d['xan_order'] for e, v in zip(cv[nm]['E'], cv[nm]['mu']) if x0 <= e <= x1]
    b.set_xlim(x0, x1)
    b.set_xticks(range(x0, x1 + 1, d['xstep']))
    b.set_ylim(-0.05, max(vals) * (1.25 + 0.07 * len(d['xan_order'])))
    b.set_yticks([])
    b.tick_params(top=False, right=False)
    b.set_xlabel('Energy (eV)')
    b.set_ylabel('Normalized μ(E) (a.u.)')
    b.legend(loc='upper left', frameon=False, fontsize=fs * 0.85, handlelength=1.6,
             labelspacing=0.25, borderaxespad=0.2)
    tag = key.replace('_L3', '-L3')
    for ext in ('png', 'pdf'):
        f.savefig(os.path.join(out_dir, f'{tag}_XAFS.{ext}'), dpi=600 if ext == 'png' else None)
    plt.close(f)


def write_origin(data, style, out_dir):
    tpl = open(os.path.join(HERE, 'origin_template.py'), encoding='utf-8').read()
    slim = {}
    for k, d in data.items():
        dd = dict(d)
        dd['curves'] = [{kk: c[kk] for kk in ('name', 'label', 'color', 'offset', 'ft', 'E', 'mu')}
                        for c in d['curves']]
        slim[k] = dd
    tpl = (tpl.replace('__OUT_DIR__', out_dir)
              .replace('__ELEMENTS__', repr(list(data.keys())))
              .replace('__DATA__', json.dumps(slim, separators=(',', ':'))))
    tpl = re.sub(r"FONT       = '.*?'", f"FONT       = '{style['font']}'", tpl)
    tpl = re.sub(r'FONT_SIZE  = \d+', f"FONT_SIZE  = {style['origin_font_size']}", tpl)
    tpl = re.sub(r'FRAME_W    = \d+', f"FRAME_W    = {style['origin_frame_width']}", tpl)
    tpl = re.sub(r'LINE_W     = \d+', f"LINE_W     = {style['origin_line_width']}", tpl)
    open(os.path.join(out_dir, 'XAFS_origin_plot.py'), 'w', encoding='utf-8').write(tpl)


def write_excel(sheets, out_dir):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print('  （未安装 openpyxl，跳过 Excel 输出：pip install openpyxl）')
        return
    wb = Workbook()
    ws0 = wb.active
    ws0.title = '说明'
    for i, t in enumerate([
        '每个数据表前 3 行：Long Name / Units / Comments，数据从第 4 行开始。',
        'Origin 导入：Data → Import From File → Excel；Column Labels 设 Long Name=1、Units=2、Comments=3，数据起始行=4。',
        '*_FT：R 之后依次为 原始 |χ(R)|、归一化（最大值=1）。*_XANES：X/Y 成对（Energy, 归一化 μ(E)）。',
    ], 1):
        ws0.cell(i, 1, t)
    ws0.column_dimensions['A'].width = 120
    F = Font(name='Arial', bold=True)
    H = PatternFill('solid', fgColor='DDEBF7')

    def col(ws, j, ln, un, cm, data):
        for r, v in enumerate([ln, un, cm], 1):
            c = ws.cell(r, j, v)
            c.font, c.fill = F, H
        for r, v in enumerate(data, 4):
            ws.cell(r, j, float(v))
    for key, d in sheets:
        tag = key.replace('_L3', '-L3')
        ws = wb.create_sheet(f'{tag}_FT')
        col(ws, 1, 'R', 'Å', '', d['R'])
        j = 2
        for c in d['curves']:
            col(ws, j, '|χ(R)|', 'Å^-(k+1)', c['name'], c['ft_raw'])
            j += 1
        for c in d['curves']:
            col(ws, j, 'FT normalized', 'a.u.', c['name'], c['ft'])
            j += 1
        ws = wb.create_sheet(f'{tag}_XANES')
        j = 1
        for nm in d['xan_order']:
            c = next(c for c in d['curves'] if c['name'] == nm)
            col(ws, j, 'Energy', 'eV', nm, c['E'])
            col(ws, j + 1, 'Normalized μ(E)', 'a.u.', nm, c['mu'])
            j += 2
    wb.save(os.path.join(out_dir, 'XAFS_data.xlsx'))


if __name__ == '__main__':
    main()
